import os
import io
import tempfile
from datetime import datetime
from typing import Dict, List, Optional

import pandas as pd
import plotly.io as pio

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

from matplotlib.figure import Figure
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import numpy as np


class ReportGenerator:
    def __init__(self):
        self._title_font_color = RGBColor(0x1E, 0x3A, 0x5F)
        self._accent_color = RGBColor(0x3B, 0x82, 0xF6)
        self._danger_color = RGBColor(0xEF, 0x44, 0x44)
        self._success_color = RGBColor(0x10, 0xB9, 0x81)

    def generate_report(
        self,
        aggregation_data: Dict,
        anomalies: List[Dict],
        forecast_data: Optional[Dict],
        forecast_alerts: List[Dict],
        dependency_metrics: Dict,
        output_path: Optional[str] = None,
    ) -> bytes:
        if not PPTX_AVAILABLE:
            raise ImportError("python-pptx is required for report generation")

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        self._add_title_slide(prs)
        self._add_summary_slide(prs, aggregation_data, dependency_metrics)
        self._add_timeseries_slide(prs, aggregation_data)
        self._add_api_metrics_slide(prs, aggregation_data)
        self._add_service_metrics_slide(prs, aggregation_data)
        self._add_anomalies_slide(prs, anomalies)
        self._add_forecast_slide(prs, forecast_data, forecast_alerts)
        self._add_dependency_slide(prs, dependency_metrics)

        if output_path:
            prs.save(output_path)
            return open(output_path, "rb").read()
        else:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as f:
                temp_path = f.name
            prs.save(temp_path)
            with open(temp_path, "rb") as f:
                data = f.read()
            os.unlink(temp_path)
            return data

    def _add_title_slide(self, prs: "Presentation"):
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        bg.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "服务器日志监控报告"
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(1))
        tf = sub_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)

        info_box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(11), Inches(1))
        tf = info_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "System Log Analytics Dashboard - 企业级实时监控与智能预警平台"
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    def _add_summary_slide(self, prs: "Presentation", agg_data: Dict, dep_metrics: Dict):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        self._add_slide_header(slide, "监控概览")

        total_logs = agg_data.get("total_logs", 0)
        error_logs = agg_data.get("error_logs", 0)
        window = agg_data.get("window_size", 0)
        error_rate = (error_logs / total_logs * 100) if total_logs > 0 else 0

        stats = [
            ("总日志数", f"{total_logs:,}", self._accent_color),
            ("错误日志数", f"{error_logs:,}", self._danger_color if error_logs > 0 else self._success_color),
            ("错误率", f"{error_rate:.2f}%", self._danger_color if error_rate > 1 else self._success_color),
            ("滑动窗口", f"{window // 60} 分钟", self._title_font_color),
            ("服务节点数", str(dep_metrics.get("service_count", 0)), self._accent_color),
            ("调用链记录", f"{dep_metrics.get('traces_tracked', 0):,}", self._title_font_color),
        ]

        cols = 3
        rows = 2
        box_width = Inches(3.8)
        box_height = Inches(1.8)
        start_x = Inches(0.7)
        start_y = Inches(2.2)
        gap_x = Inches(0.2)
        gap_y = Inches(0.3)

        for idx, (label, value, color) in enumerate(stats):
            row = idx // cols
            col = idx % cols
            x = start_x + col * (box_width + gap_x)
            y = start_y + row * (box_height + gap_y)

            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_width, box_height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
            shape.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)

            tf = shape.text_frame
            tf.margin_left = Inches(0.3)
            tf.margin_top = Inches(0.2)

            p1 = tf.paragraphs[0]
            p1.text = label
            p1.font.size = Pt(16)
            p1.font.color.rgb = self._title_font_color
            p1.font.bold = True

            p2 = tf.add_paragraph()
            p2.text = value
            p2.font.size = Pt(32)
            p2.font.bold = True
            p2.font.color.rgb = color

    def _add_timeseries_slide(self, prs: "Presentation", agg_data: Dict):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_slide_header(slide, "请求量与响应时间趋势")

        ts = agg_data.get("timeseries", {})
        timestamps = ts.get("timestamps", [])
        req_counts = ts.get("request_counts", [])
        avg_rts = ts.get("avg_response_times", [])

        if not timestamps:
            self._add_empty_note(slide, "暂无趋势数据")
            return

        fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(11.5, 5), gridspec_kw={"height_ratios": [2, 1]})

        try:
            dt_objects = [pd.to_datetime(t) for t in timestamps]
        except Exception:
            dt_objects = list(range(len(timestamps)))

        ax1.bar(range(len(req_counts)), req_counts, color="#3B82F6", alpha=0.7)
        ax1.set_ylabel("请求数 (count/10s)")
        ax1.set_title("请求量时序")
        ax1.grid(True, alpha=0.3)

        if isinstance(dt_objects[0], datetime) or hasattr(dt_objects[0], "hour"):
            ax1.set_xticks(range(0, len(dt_objects), max(1, len(dt_objects) // 6)))
            ax1.set_xticklabels([dt_objects[i].strftime("%H:%M:%S") if isinstance(dt_objects[i], datetime) else str(dt_objects[i])
                                 for i in range(0, len(dt_objects), max(1, len(dt_objects) // 6))], fontsize=8)

        ax2.plot(range(len(avg_rts)), avg_rts, color="#F59E0B", linewidth=1.5)
        ax2.set_ylabel("平均响应 (ms)")
        ax2.set_title("平均响应时间")
        ax2.grid(True, alpha=0.3)
        if isinstance(dt_objects[0], datetime) or hasattr(dt_objects[0], "hour"):
            ax2.set_xticks(range(0, len(dt_objects), max(1, len(dt_objects) // 6)))
            ax2.set_xticklabels([dt_objects[i].strftime("%H:%M:%S") if isinstance(dt_objects[i], datetime) else str(dt_objects[i])
                                 for i in range(0, len(dt_objects), max(1, len(dt_objects) // 6))], fontsize=8)

        plt.tight_layout()

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            img_path = f.name
        fig.savefig(img_path, dpi=150, bbox_inches="tight")
        plt.close(fig)

        slide.shapes.add_picture(img_path, Inches(0.9), Inches(1.8), width=Inches(11.5))
        os.unlink(img_path)

    def _add_api_metrics_slide(self, prs: "Presentation", agg_data: Dict):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_slide_header(slide, "API 接口性能排行 (Top 10)")

        api_metrics = agg_data.get("api_metrics", {})
        if not api_metrics:
            self._add_empty_note(slide, "暂无API数据")
            return

        api_list = [
            {"name": name, **metrics}
            for name, metrics in api_metrics.items()
        ]
        api_list.sort(key=lambda x: x.get("total_requests", 0), reverse=True)
        api_list = api_list[:10]

        fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))

        names = [a["name"].split("/")[-1] for a in api_list]
        error_rates = [a.get("error_rate", 0) * 100 for a in api_list]
        p99_rts = [a.get("p99", 0) for a in api_list]

        colors = ["#EF4444" if e > 5 else "#F59E0B" if e > 1 else "#10B981" for e in error_rates]
        ax1.barh(range(len(names)), error_rates, color=colors)
        ax1.set_yticks(range(len(names)))
        ax1.set_yticklabels(names, fontsize=8)
        ax1.set_xlabel("错误率 (%)")
        ax1.set_title("API 错误率")
        ax1.invert_yaxis()
        ax1.grid(True, alpha=0.3, axis="x")

        ax2.barh(range(len(names)), p99_rts, color="#3B82F6")
        ax2.set_yticks(range(len(names)))
        ax2.set_yticklabels(names, fontsize=8)
        ax2.set_xlabel("P99 响应时间 (ms)")
        ax2.set_title("API P99 响应时间")
        ax2.invert_yaxis()
        ax2.grid(True, alpha=0.3, axis="x")

        plt.tight_layout()

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            img_path = f.name
        fig.savefig(img_path, dpi=150, bbox_inches="tight")
        plt.close(fig)

        slide.shapes.add_picture(img_path, Inches(0.6), Inches(1.8), width=Inches(12))
        os.unlink(img_path)

    def _add_service_metrics_slide(self, prs: "Presentation", agg_data: Dict):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_slide_header(slide, "服务节点健康状态")

        service_metrics = agg_data.get("service_metrics", {})
        if not service_metrics:
            self._add_empty_note(slide, "暂无服务数据")
            return

        data = []
        for service, metrics in service_metrics.items():
            data.append({
                "服务": service,
                "请求数": metrics.get("total_requests", 0),
                "错误数": metrics.get("error_count", 0),
                "错误率(%)": round(metrics.get("error_rate", 0) * 100, 2),
                "P99响应(ms)": metrics.get("p99_rt", 0),
            })

        df = pd.DataFrame(data).sort_values("请求数", ascending=False)
        table = self._create_table_from_df(slide, df, Inches(0.6), Inches(2), Inches(12), Inches(4.5))

    def _add_anomalies_slide(self, prs: "Presentation", anomalies: List[Dict]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_slide_header(slide, f"异常检测事件 (共 {len(anomalies)} 条)")

        if not anomalies:
            self._add_empty_note(slide, "无异常事件 - 系统运行正常")
            return

        data = []
        for a in anomalies[:15]:
            data.append({
                "时间": a.get("timestamp", "")[-8:],
                "严重程度": a.get("severity", ""),
                "类型": a.get("type", ""),
                "实体": a.get("service", a.get("endpoint", "")),
                "指标": a.get("metric", ""),
                "当前值": a.get("value", ""),
                "阈值": a.get("threshold", ""),
                "Z分数": a.get("z_score", "-"),
            })

        df = pd.DataFrame(data)
        self._create_table_from_df(slide, df, Inches(0.3), Inches(1.8), Inches(12.7), Inches(5))

    def _add_forecast_slide(self, prs: "Presentation", forecast_data: Optional[Dict], alerts: List[Dict]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_slide_header(slide, "流量预测与预警")

        if not forecast_data:
            self._add_empty_note(slide, "暂无预测数据")
            return

        timestamps = forecast_data.get("timestamps", [])
        predicted = forecast_data.get("predicted", [])
        lower = forecast_data.get("lower_bound", [])
        upper = forecast_data.get("upper_bound", [])
        threshold = forecast_data.get("threshold", 0)

        fig, ax = plt.subplots(figsize=(11.5, 4.5))

        x = range(len(predicted))
        ax.plot(x, predicted, color="#3B82F6", label="预测值", linewidth=2)
        ax.fill_between(x, lower, upper, alpha=0.2, color="#3B82F6", label="置信区间")
        ax.axhline(y=threshold, color="#EF4444", linestyle="--", label=f"阈值 ({threshold:,.0f})")
        ax.set_xlabel("预测时间 (1h / 30s 粒度)")
        ax.set_ylabel("请求数")
        ax.set_title(f"未来1小时流量预测 ({forecast_data.get('model', '')} 模型)")
        ax.legend()
        ax.grid(True, alpha=0.3)

        xticks = list(range(0, len(x), len(x) // 6))
        ax.set_xticks(xticks)
        ax.set_xticklabels([timestamps[i][-8:] for i in xticks], fontsize=8, rotation=30)

        plt.tight_layout()

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            img_path = f.name
        fig.savefig(img_path, dpi=150, bbox_inches="tight")
        plt.close(fig)

        slide.shapes.add_picture(img_path, Inches(0.9), Inches(1.8), width=Inches(11.5))

        if alerts:
            alert_text = "\n".join([f"• [{a.get('severity','')}] {a.get('message','')}" for a in alerts[:3]])
            alert_box = slide.shapes.add_textbox(Inches(0.9), Inches(6.2), Inches(11.5), Inches(1))
            tf = alert_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = "预测预警:"
            p.font.bold = True
            p.font.size = Pt(14)
            p.font.color.rgb = self._danger_color
            for line in alert_text.split("\n"):
                pp = tf.add_paragraph()
                pp.text = line
                pp.font.size = Pt(11)
                pp.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    def _add_dependency_slide(self, prs: "Presentation", dep_metrics: Dict):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_slide_header(slide, "服务调用依赖概览")

        stats = [
            ("服务节点数", str(dep_metrics.get("service_count", 0))),
            ("调用连边数", str(dep_metrics.get("edge_count", 0))),
            ("总调用次数", f"{dep_metrics.get('total_calls', 0):,}"),
            ("总错误数", f"{dep_metrics.get('total_errors', 0):,}"),
            ("平均响应(ms)", str(dep_metrics.get("average_response_time_ms", 0))),
            ("跟踪的调用链", f"{dep_metrics.get('traces_tracked', 0):,}"),
        ]

        cols = 3
        rows = 2
        box_width = Inches(3.8)
        box_height = Inches(1.8)
        start_x = Inches(0.7)
        start_y = Inches(2.2)
        gap_x = Inches(0.2)
        gap_y = Inches(0.3)

        for idx, (label, value) in enumerate(stats):
            row = idx // cols
            col = idx % cols
            x = start_x + col * (box_width + gap_x)
            y = start_y + row * (box_height + gap_y)

            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_width, box_height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
            shape.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

            tf = shape.text_frame
            tf.margin_left = Inches(0.3)
            tf.margin_top = Inches(0.2)

            p1 = tf.paragraphs[0]
            p1.text = label
            p1.font.size = Pt(16)
            p1.font.color.rgb = self._title_font_color
            p1.font.bold = True

            p2 = tf.add_paragraph()
            p2.text = value
            p2.font.size = Pt(28)
            p2.font.bold = True
            p2.font.color.rgb = self._accent_color

    def _add_slide_header(self, slide, title: str):
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.2))
        header.fill.solid()
        header.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        header.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        footer = slide.shapes.add_textbox(Inches(0.6), Inches(7.1), Inches(12), Inches(0.3))
        tf = footer.text_frame
        p = tf.paragraphs[0]
        p.text = "System Log Analytics Dashboard - 企业级实时监控平台"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    def _add_empty_note(self, slide, msg: str):
        box = slide.shapes.add_textbox(Inches(4), Inches(3), Inches(5), Inches(1))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = 1
        p.text = msg
        p.font.size = Pt(20)
        p.font.color.rgb = self._accent_color

    def _create_table_from_df(self, slide, df: pd.DataFrame, x, y, width, height):
        rows, cols = df.shape
        rows += 1

        table_shape = slide.shapes.add_table(rows, cols, x, y, width, height)
        table = table_shape.table

        for col_idx, col_name in enumerate(df.columns):
            cell = table.cell(0, col_idx)
            cell.text = str(col_name)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(12)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self._title_font_color

        for row_idx in range(df.shape[0]):
            for col_idx in range(df.shape[1]):
                cell = table.cell(row_idx + 1, col_idx)
                value = df.iloc[row_idx, col_idx]
                cell.text = str(value)
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                if row_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)

        return table
